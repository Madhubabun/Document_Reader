"""Builds the small Office files used by the OOXML reader tests.

Run: pip install python-docx openpyxl python-pptx && python3 tool/make_fixtures.py
The outputs are committed under test/fixtures/ so tests need no Python.
"""
import io
import struct
import zlib
from pathlib import Path

import docx
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent.parent / "test" / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)


def tiny_png(rgb=(255, 154, 60), size=4):
    row = b"\x00" + bytes(rgb) * size
    raw = row * size

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def make_docx():
    d = docx.Document()
    d.add_heading("Project Brief", level=0)
    d.add_heading("Objectives", level=1)
    p = d.add_paragraph("Plain start, ")
    b = p.add_run("bold part")
    b.bold = True
    i = p.add_run(" and italic red")
    i.italic = True
    i.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    i.font.size = Pt(14)
    d.add_paragraph("First bullet", style="List Bullet")
    d.add_paragraph("Second bullet", style="List Bullet")
    c = d.add_paragraph("Centered line")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Phase"
    t.cell(0, 1).text = "Weeks"
    t.cell(1, 0).text = "Survey"
    t.cell(1, 1).text = "1-2"
    d.add_picture(io.BytesIO(tiny_png()), width=Inches(1))
    d.save(OUT / "sample.docx")


def make_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Item"
    ws["B1"] = "Cost"
    ws["A2"] = "Laptop"
    ws["B2"] = 1200
    ws["A3"] = "Coffee"
    ws["B3"] = 0.1 + 0.2
    ws["B4"] = "=SUM(B2:B3)"
    ws["D10"] = True
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Shared string again: Laptop"
    wb.save(OUT / "sample.xlsx")


def make_pptx():
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Q3 Highlights"
    s1.placeholders[1].text = "Pitch deck"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Results"
    body = s2.placeholders[1].text_frame
    body.text = "Revenue up"
    para = body.add_paragraph()
    para.text = "New markets"
    para.level = 1
    s2.shapes.add_picture(io.BytesIO(tiny_png((76, 141, 255))), Inches(9), Inches(2), Inches(3), Inches(2))
    box = s2.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Footnote"
    run.font.bold = True
    run.font.size = Pt(12)
    prs.save(OUT / "sample.pptx")


make_docx()
make_xlsx()
make_pptx()
print("fixtures written to", OUT)
